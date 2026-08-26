import openpyxl, json, sys, os, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

desktop = os.path.join(os.path.expanduser("~"), "Desktop")
files = [f for f in os.listdir(desktop) if f.startswith("基金持仓")]
if files:
    path = os.path.join(desktop, files[0])
    print(f"Reading: {path}")
    wb = openpyxl.load_workbook(path, data_only=True)
    for sn in wb.sheetnames:
        ws = wb[sn]
        print(f"\n=== {sn} ===")
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = [str(v)[:50] if v is not None else "" for v in row]
            while vals and vals[-1] == "":
                vals.pop()
            if any(v != "" for v in vals):
                print(f"R{i}: {vals}")
            if i > 30:
                break
